from openpyxl import Workbook
from openpyxl.styles import Font
from typing import Dict, Any


def _write_dict_to_sheet(ws, data: Dict[str, Any], field_order: list[str]):
    
    ws.cell(row=1, column=1, value="Field")
    ws.cell(row=1, column=2, value="Value")
    
    bold_font = Font(bold=True)
    ws["A1"].font = bold_font
    ws["B1"].font = bold_font
    
    row = 2
    for field in field_order:
        label = field.replace("_", " ").capitalize()
        value = data.get(field, "")
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=str(
            value) if value is not None else "")
        row += 1


def export_plan_to_excel(plan_data: Dict[str, Any], field_order: list[str]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"  # type: ignore

    _write_dict_to_sheet(ws, plan_data, field_order)
    return wb


def export_multiple_plans_to_excel(plans: list[Dict[str, Any]], field_order: list[str]) -> Workbook:
    wb = Workbook()
    first = True

    for idx, plan in enumerate(plans, start=1):
        if first:
            ws = wb.active
            ws.title = f"Plan_{idx}"  # type: ignore
            first = False
        else:
            ws = wb.create_sheet(title=f"Plan_{idx}")

        _write_dict_to_sheet(ws, plan, field_order)

    return wb
