from typing import List, Dict
from openpyxl import Workbook
from openpyxl.styles import Font
from app.models.requirement import Requirement
from app.models.defect import Defect

EXPORT_COLUMNS = [
    ("Defect ID", "defect_id"),
    ("Module", "module"),
    ("Description", "description"),
    ("Status", "status"),
    ("Bug Detected Date", "bug_detected_date"),
    ("Fixed Date", "fixed_date"),
    ("Reopen Date", "reopen_date"),
    ("Remarks", "remarks"),
]


def export_defects_to_excel(requirement_name: str, defects: List[Defect]) -> Workbook:
    wb = Workbook()
    if wb.active is None:
        ws = wb.create_sheet(title=requirement_name[:31])
    else:
        ws = wb.active
        ws.title = requirement_name[:31]

    # Header
    for col_idx, (header, _) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    row_idx = 2
    for defect in defects:
        if defect.is_deleted:
            continue
        data: Dict = defect.defect_data or {}
        for col_idx, (_, key) in enumerate(EXPORT_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=data.get(key, ""))
        row_idx += 1

    return wb


def export_project_defects_to_excel(requirements: List[Requirement]) -> Workbook:
    wb = Workbook()

    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    for req in requirements:
        sheet_name = str(req.req_id)[:31]  # Excel sheet name max 31 chars
        ws = wb.create_sheet(title=sheet_name)

        # Write header
        for col_idx, (header, _) in enumerate(EXPORT_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True)

        # Write defect rows
        row_idx = 2
        for defect in req.defects:
            if defect.is_deleted:
                continue
            data: Dict = defect.defect_data or {}
            for col_idx, (_, key) in enumerate(EXPORT_COLUMNS, start=1):
                ws.cell(row=row_idx, column=col_idx, value=data.get(key, ""))
            row_idx += 1

    return wb
