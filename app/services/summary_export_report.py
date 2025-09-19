import io
import json
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from typing import List
from app.models.summary_report import SummaryReport


def export_summary_reports_to_excel(reports: List[SummaryReport], project_name: str) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet(title="Summary Report")
    else:
        ws.title = "Summary Report"

    headers = [
        "S NO", "Testing Features", "No of test items", "No of test implements",
        "No of test non-implements", "OK", "NOT OKAY", "Test Coverage (%)"
    ]

    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header).font = Font(bold=True)

    current_row = 2
    for report in reports:
        # Ensure report.data is a list of dicts
        raw_data = report.data or {}
        
        if isinstance(raw_data, str):
            try:
                raw_data = json.loads(raw_data)
            except json.JSONDecodeError:
                raw_data = {}

        if isinstance(raw_data, dict):
            data = raw_data.get("summary_data", [])
        else:
            data = raw_data

        if not isinstance(data, list):
            data = []

        for i, feature in enumerate(data, start=1):
            if not isinstance(feature, dict):
                continue
            ws.cell(row=current_row, column=1, value=i)
            ws.cell(row=current_row, column=2,
                    value=feature.get("feature_name", ""))
            ws.cell(row=current_row, column=3,
                    value=feature.get("total_items", 0))
            ws.cell(row=current_row, column=4,
                    value=feature.get("implemented", 0))
            ws.cell(row=current_row, column=5,
                    value=feature.get("non_implemented", 0))
            ws.cell(row=current_row, column=6, value=feature.get("ok", 0))
            ws.cell(row=current_row, column=7, value=feature.get("not_ok", 0))
            ws.cell(row=current_row, column=8,
                    value=feature.get("coverage", 0))
            current_row += 1

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    project_name_safe = project_name.replace(" ", "_")[:50]
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={project_name_safe}_summary_report.xlsx"}
    )
