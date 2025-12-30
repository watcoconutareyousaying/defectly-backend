import io
from typing import List, Dict
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell
from app.models.case import Case
from app.models.requirement import Requirement

EXPORT_COLUMNS = [
    ("Test Case ID", "case_id"),
    ("Module", "module"),
    ("Title", "title"),
    ("Description", "description"),
    ("Test Steps", "steps"),
    ("Test Data", "data"),
    ("Expected Result", "expected_result"),
    ("Priority", "priority"),
    ("Status", "status"),
    ("Remarks", "remarks"),
]


def export_cases_to_excel(requirement_name: str, cases: List[Case]) -> Workbook:
    wb = Workbook()
    ws: Worksheet
    if wb.active is None:
        ws = wb.create_sheet(title=requirement_name[:31])
    else:
        ws = wb.active
        ws.title = requirement_name[:31]

    # Write header
    for col_idx, (header, _) in enumerate(EXPORT_COLUMNS, start=1):
        cell: Cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    # Write case rows
    row_idx = 2
    for case in cases:
        if case.is_deleted:
            continue
        data: Dict = case.case_data or {}
        for col_idx, (_, key) in enumerate(EXPORT_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=data.get(key, ""))
        row_idx += 1

    return wb


# Multi-requirement export (project export)
def export_project_cases_to_excel(requirements: List[Requirement]) -> Workbook:
    wb = Workbook()

    # Remove default sheet if exists
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    for req in requirements:
        sheet_name = str(req.req_id)[:31]  # Excel sheet name max 31 chars
        ws: Worksheet = wb.create_sheet(title=sheet_name)

        # Header row
        for col_idx, (header, _) in enumerate(EXPORT_COLUMNS, start=1):
            cell: Cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True)

        # Case rows
        row_idx = 2
        for case in req.cases:
            if case.is_deleted:
                continue
            data: Dict = case.case_data or {}
            for col_idx, (_, key) in enumerate(EXPORT_COLUMNS, start=1):
                ws.cell(row=row_idx, column=col_idx, value=data.get(key, ""))
            row_idx += 1

    return wb
