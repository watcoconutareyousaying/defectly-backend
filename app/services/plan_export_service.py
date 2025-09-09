import io
import re
from typing import Dict, Any, List
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session
from app.services import plan_service


FIELD_ORDER = [
    "project_name",
    "version",
    "prepared_by",
    "date",
    "module",
    "phase",
    "objective",
    "scope",
    "test_environment",
    "tools_used",
    "entry_criteria",
    "exit_criteria",
    "risks_and_mitigations",
]


def _sanitize_filename(name: str) -> str:
    safe = re.sub(r"[^a-zA-Z0-9]+", "_", name.strip().lower())
    return safe


def _write_dict_to_sheet(ws, data: Dict[str, Any], field_order: List[str]):
    ws.cell(row=1, column=1, value="Field")
    ws.cell(row=1, column=2, value="Value")

    bold_font = Font(bold=True)
    ws["A1"].font = bold_font
    ws["B1"].font = bold_font

    row = 2
    for field in field_order:
        label = field.replace("_", " ").capitalize()
        value = data.get(field, "")
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=str(value) if value is not None else "")
        row += 1


def _export_plan_to_excel(plan_data: Dict[str, Any], field_order: List[str]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    if ws is not None:
        ws.title = "Plan"
        _write_dict_to_sheet(ws, plan_data, field_order)
    return wb


def _export_multiple_plans_to_excel(plans: List[Dict[str, Any]], field_order: List[str]) -> Workbook:
    wb = Workbook()
    for idx, plan in enumerate(plans, start=1):
        if idx == 1:
            ws = wb.active
            if ws is not None:
                ws.title = f"Plan_{idx}"
        else:
            ws = wb.create_sheet(title=f"Plan_{idx}")
        _write_dict_to_sheet(ws, plan, field_order)
    return wb


def export_single_plan(db: Session, plan_id: int, user_id: int):
    plan = plan_service.get_plan_by_id(db, plan_id)
    plan_dict = {
        field: plan.plan_data.get(field) or (plan.creator.name if field == "prepared_by" else None)
        for field in FIELD_ORDER
    }

    wb = _export_plan_to_excel(plan_dict, FIELD_ORDER)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"{_sanitize_filename(plan.plan_data.get('project_name'))}-testplan.xlsx"

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


def export_project_plans(db: Session, project_id: int):
    plans = plan_service.list_plans_for_project(db, project_id)

    plan_dicts = [
        {field: plan.plan_data.get(field) or (plan.creator.name if field == "prepared_by" else None)
         for field in FIELD_ORDER}
        for plan in plans
    ]

    wb = _export_multiple_plans_to_excel(plan_dicts, FIELD_ORDER)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    project_name = plans[0].plan_data.get("project_name", f"project_{project_id}")
    filename = f"{_sanitize_filename(project_name)}-testplan.xlsx"

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
